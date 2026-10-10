"""Import a club's existing Excel scoring workbook (boats list + race sheets).

Layout it understands (the common "one sheet per race" workbook):
  * "Boats List": name, design, sail #, skipper, club, rating, type flag (S/NS/O), address, phone, email
    (rows named "X W/L" and "X R/L" become one boat with two ratings)
  * "Race #N": one block per class: a "Course = n  Distance = d  START hh:mm:ss" row, then boat rows
  * "Series Standings": boat names with y/n in the "Entered Series" column

Every imported finish is checked against the sheet's own CORRECTED column and any difference is printed.
"""
import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from boats.models import Boat, BoatRating, Fleet, RatingType
from racing import services
from racing.models import Finish, Payment, Race, Series, SeriesEntry, Start

STATUS = {"dns", "dnf", "dsq", "ocs", "ret", "dnc", "tle", "dne", "ufd", "bfd", "rcd", "rdg"}
SUFFIX = re.compile(r"\s+[RW]/(?:L)?$", re.I)


def norm(name: str) -> str:
    return re.sub(r"\s+", " ", SUFFIX.sub("", str(name or "").strip())).lower()


def fleet_name(cls: str) -> str:
    return re.sub(r"\s*-\s*", "-", re.sub(r"\s+", " ", cls.strip())).upper()


def secs(v) -> int | None:
    if isinstance(v, time):
        return v.hour * 3600 + v.minute * 60 + v.second + (1 if v.microsecond >= 500_000 else 0)
    if isinstance(v, timedelta):
        return int(round(v.total_seconds()))
    return None


class Command(BaseCommand):
    help = "Import boats, ratings and race results from a club Excel workbook"

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--races", default="1", help='Race numbers to import, e.g. "1" or "1,2". Default: 1')
        parser.add_argument("--rating-type", default="RL", help="Rating type for PHRF blocks (default RL)")
        parser.add_argument("--ratings-effective", default=None, help="YYYY-MM-DD for boat ratings (default Jan 1)")
        parser.add_argument("--boats-only", action="store_true")
        parser.add_argument("--member-club", default=None,
                            help="Home-club code whose boats are members (default: CLUB_SHORT_NAME)")

    @transaction.atomic
    def handle(self, *args, **o):
        try:
            import openpyxl
        except ImportError:
            raise CommandError("pip install openpyxl (or: uv sync) first")
        wb = openpyxl.load_workbook(o["path"], data_only=True)
        types = {t.code: t for t in RatingType.objects.all()}
        for need in ("WL", "RL", "ORCA"):
            if need not in types:
                raise CommandError(f"Rating type {need} missing; run migrate first")
        series_year = int(wb["Race #1"]["A1"].value or date.today().year) if "Race #1" in wb.sheetnames else date.today().year
        eff = date.fromisoformat(o["ratings_effective"]) if o["ratings_effective"] else date(series_year, 1, 1)
        from django.conf import settings
        member_club = (o["member_club"] or settings.CLUB_SHORT_NAME or "").strip().lower()
        boats = self._boats(wb["Boats List"], types, eff)
        if member_club:
            n = Boat.objects.filter(home_club__iexact=member_club).update(is_member=True)
            self.stdout.write(f"Members: {n} boats from club {member_club.upper()} marked as members")
        self.stdout.write(f"Boats: {len(boats)} imported/updated")
        if o["boats_only"]:
            return
        for n in [int(x) for x in o["races"].split(",") if x.strip()]:
            if f"Race #{n}" not in wb.sheetnames:
                raise CommandError(f"No sheet 'Race #{n}'")
            self._race(wb, n, boats, types, o["rating_type"])
        if "Accounting" in wb.sheetnames:
            self._payments(wb["Accounting"], boats, series_year)

    # ------------------------------------------------------------------
    def _boats(self, ws, types, eff):
        boats: dict[str, Boat] = {}
        for row in ws.iter_rows(values_only=True):
            name, flag = row[0], str(row[6] or "").strip().upper()
            if not name or str(name).strip() in ("PARTIAL INFO", "FULL INFO FOUND") or flag not in ("S", "NS", "O"):
                continue
            raw = str(name).strip()
            key = norm(raw)
            suffix = re.search(r"([RW])/", raw[len(SUFFIX.sub("", raw)):], re.I)
            try:
                rating = Decimal(str(row[5])) if row[5] not in (None, "") else None
            except Exception:
                rating = None
            boat = boats.get(key)
            if boat is None:
                sail = str(row[2] or "").strip()[:20]
                boat, _ = Boat.objects.update_or_create(
                    name=SUFFIX.sub("", raw).strip()[:80], sail_number=sail,
                    defaults={"design": str(row[1] or "")[:60], "skipper": str(row[3] or "")[:80],
                              "home_club": str(row[4] or "")[:20],
                              "contact_phone": str(row[8] or "").replace("\xa0", "").strip()[:30],
                              "contact_email": (str(row[9] or "").split(",")[0].strip() if "@" in str(row[9] or "") else "")})
                boats[key] = boat
            if flag == "O" and boat.fixed_rating_type_id is None:
                boat.fixed_rating_type = types["ORCA"]          # ORCA boats always use ORCA
                boat.save(update_fields=["fixed_rating_type"])
            if rating is None:
                continue
            if flag == "O":
                codes = ["ORCA"]
            elif suffix:
                codes = ["WL" if suffix.group(1).upper() == "W" else "RL"]
            else:
                codes = ["WL", "RL"]       # one rating given: use it for both until the club splits them
            for code in codes:
                BoatRating.objects.update_or_create(
                    boat=boat, rating_type=types[code], effective_from=eff, defaults={"value": rating})
        return boats

    # ------------------------------------------------------------------
    def _race(self, wb, number, boats, types, phrf_code):
        ws = wb[f"Race #{number}"]
        series_name, year = str(ws["B1"].value).strip(), int(ws["A1"].value)
        race_date = ws["I1"].value.date()
        series, _ = Series.objects.get_or_create(
            name=series_name, season=year, defaults={"status": Series.Status.ACTIVE, "scoring_basis": "entries"})
        race, _ = Race.objects.get_or_create(series=series, number=number, defaults={"date": race_date})
        race.date, race.status = race_date, Race.Status.SAILED
        race.save()
        entered = self._entered(wb)
        rows = list(ws.iter_rows(values_only=True))
        blocks, cur = [], None
        for r in rows:
            if str(r[3] or "").strip().startswith("Course"):
                cur = {"head": r, "boats": [], "cls": None}
                blocks.append(cur)
            elif cur is not None and r[1] and str(r[1]).strip() and str(r[0] or "") != "CLASS":
                if r[0] and not cur["cls"]:
                    cur["cls"] = str(r[0])
                cur["boats"].append(r)
            elif cur is not None and r[0] and str(r[0]) != "CLASS" and not cur["cls"]:
                cur["cls"] = str(r[0])
        mism = total = 0
        for b in blocks:
            if not b["boats"]:
                continue
            h = b["head"]
            fleet, _ = Fleet.objects.get_or_create(name=fleet_name(b["cls"] or "FLEET"))
            series.fleets.add(fleet)
            orca = "ORCA" in str(h[11] or "").upper()
            rt = types["ORCA" if orca else phrf_code]
            start_t = h[9] if isinstance(h[9], time) else time(0, 0)
            course_label = f"{h[3]} {h[4]}".replace("Course =", "").strip()
            dist = Decimal(str(h[6] or 0)) or None
            start, _ = Start.objects.update_or_create(
                race=race, fleet=fleet,
                defaults={"start_time": datetime.combine(race_date, start_t), "rating_type": rt,
                          "distance_override": dist, "course": None})
            self.stdout.write(f"  {fleet.name}: start {start_t}, {dist} nm, {rt.code} (course {course_label})")
            for r in b["boats"]:
                key = norm(r[1])
                boat = boats.get(key) or Boat.objects.filter(name__iexact=SUFFIX.sub("", str(r[1]).strip())).first()
                if boat is None:
                    boat = Boat.objects.create(name=SUFFIX.sub("", str(r[1]).strip())[:80], sail_number=str(r[3] or "")[:20],
                                               design=str(r[2] or "")[:60], skipper=str(r[4] or "")[:80],
                                               home_club=str(r[5] or "")[:20])
                    boats[key] = boat
                if boat.default_fleet_id is None:
                    boat.default_fleet = fleet
                    boat.save(update_fields=["default_fleet"])
                SeriesEntry.objects.update_or_create(series=series, boat=boat, defaults={
                    "fleet": fleet, "in_series": entered.get(key, False)})
                f, _ = Finish.objects.get_or_create(start=start, boat=boat, defaults={"rating_type": rt})
                fin = r[10]
                if isinstance(fin, (time, timedelta)):
                    t = datetime.combine(race_date, time.min) + (
                        fin if isinstance(fin, timedelta) else
                        timedelta(hours=fin.hour, minutes=fin.minute, seconds=fin.second))
                    f.finish_time = t if t >= start.start_time else t + timedelta(days=1)
                    f.status = ""
                elif isinstance(fin, str) and fin.strip().lower() in STATUS:
                    f.finish_time, f.status = None, fin.strip().upper()
                else:
                    continue
                try:
                    rating = Decimal(str(r[6]))
                except Exception:
                    rating = None
                services.save_finish(f, manual_value=rating)
                want = secs(r[12])
                if f.corrected_seconds is not None and want is not None:
                    total += 1
                    if abs(f.corrected_seconds - want) > 1:
                        mism += 1
                        self.stdout.write(self.style.WARNING(
                            f"    MISMATCH {boat.name}: app {f.corrected_seconds}s vs sheet {want}s"))
        self.stdout.write(self.style.SUCCESS(
            f"Race {number}: imported; {total - mism}/{total} corrected times match the sheet."))

    def _entered(self, wb):
        out = {}
        if "Series Standings" in wb.sheetnames:
            for r in wb["Series Standings"].iter_rows(values_only=True):
                if r[1] and str(r[1]).strip() not in ("0", "Boat Name") and str(r[2] or "").strip().lower() in ("y", "n"):
                    out[norm(r[1])] = str(r[2]).strip().lower() == "y"
        return out

    # ------------------------------------------------------------------
    def _payments(self, ws, boats, year):
        """Accounting sheet: a PRICES row (series, race 1, ...) and one row per boat with amounts."""
        rows = list(ws.iter_rows(values_only=True))
        prices = next((r for r in rows if str(r[7] or "").strip().upper() == "PRICES:"), None)
        hdr = next((r for r in rows if str(r[0] or "").strip() == "Boat Name"), None)
        series = Series.objects.filter(season=year).exclude(kind="single").first()
        if not (prices and hdr and series):
            return
        cols = {str(c).strip().lower().replace(" ", ""): i for i, c in enumerate(hdr) if c}
        series.series_fee = Decimal(str(prices[cols["series"]] or 0))
        race_cols = {int(k[4:]): i for k, i in cols.items() if k.startswith("race") and k[4:].isdigit()}
        series.race_fee = Decimal(str(prices[race_cols[min(race_cols)]] or 0)) if race_cols else Decimal(0)
        series.save()
        by_sail = {b.sail_number: b for b in boats.values() if b.sail_number}
        made = 0
        for r in rows[rows.index(hdr) + 1:]:
            if not r[0] or str(r[0]).strip() in ("", "0"):
                continue
            boat = by_sail.get(str(r[2] or "").strip()) or boats.get(norm(r[0]))
            if boat is None:
                self.stdout.write(self.style.WARNING(f"  payments: no boat matches {r[0]!r}"))
                continue
            items = []                       # (race number or None, item price)
            if r[cols["series"]]:
                items.append((None, Decimal(str(r[cols["series"]]))))
            for n, i in race_cols.items():
                if r[i]:
                    items.append((n, Decimal(str(r[i]))))
            paid = Decimal(str(r[cols["paid"]] or 0)) if "paid" in cols else Decimal(0)
            method = str(r[cols["howpd"]] or "")[:30] if "howpd" in cols else ""
            for n, price in items:
                race = Race.objects.filter(series=series, number=n).first() if n else None
                if n and race is None:
                    race = Race.objects.create(series=series, number=n, date=date(year, 1, 1))
                amount = paid if len(items) == 1 and paid else price
                Payment.objects.update_or_create(
                    series=series, boat=boat, race=race,
                    defaults={"amount": amount, "paid_on": date.today(), "method": method,
                              "note": "imported from the Accounting sheet"})
                made += 1
        self.stdout.write(self.style.SUCCESS(
            f"Payments: {made} imported (fees {series.series_fee} / {series.race_fee}). "
            "Payments are not enforced until 'Require payment' is turned on for the series."))
